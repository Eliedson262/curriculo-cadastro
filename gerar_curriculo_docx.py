from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_run_font(run, name='Arial', size=11, bold=False, color='000000'):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)


def add_section_title(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text.upper())
    set_run_font(r, size=11, bold=True, color='1F4E79')
    set_paragraph_spacing(p, before=8, after=4, line=1.0)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    r = p.add_run('• ' + text)
    set_run_font(r, size=11)
    set_paragraph_spacing(p, before=0, after=1, line=1.15)
    return p


def format_docx(data, output_path='curriculo_formatado.docx'):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data.get('nome', 'SEU NOME'))
    set_run_font(r, size=16, bold=True)
    set_paragraph_spacing(p, after=3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contato = []
    if data.get('telefone'):
        contato.append(data['telefone'])
    if data.get('email'):
        contato.append(data['email'])
    if data.get('cidade'):
        contato.append(data['cidade'])
    if data.get('estado'):
        contato.append(data['estado'])
    if data.get('linkedin'):
        contato.append(data['linkedin'])
    r = p.add_run(' | '.join(contato))
    set_run_font(r, size=10, color='333333')
    set_paragraph_spacing(p, after=8)

    line = doc.add_paragraph()
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '7F7F7F')
    border.append(bottom)
    line._p.get_or_add_pPr().append(border)
    set_paragraph_spacing(line, after=8)

    if data.get('objetivo'):
        add_section_title(doc, 'Objetivo profissional')
        p = doc.add_paragraph()
        r = p.add_run(data['objetivo'])
        set_run_font(r, size=11)
        set_paragraph_spacing(p, after=4, line=1.15)

    add_section_title(doc, 'Dados pessoais')
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    pairs = [
        ('Idade', data.get('idade', '')),
        ('Escolaridade', data.get('escolaridade', '')),
        ('Pretensão salarial', data.get('pretensao', '')),
        ('Localização', f"{data.get('cidade', '')}/{data.get('estado', '')}".strip('/')),
    ]
    for label, value in pairs:
        if value:
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value
            set_cell_shading(row[0], 'D9EAF7')
            for cell in row:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_font(run, size=10)

    if data.get('curso_grad') or data.get('instituicao') or data.get('ano_conclusao'):
        add_section_title(doc, 'Formação acadêmica')
        p = doc.add_paragraph()
        if data.get('curso_grad'):
            r = p.add_run(data['curso_grad'])
            set_run_font(r, size=11, bold=True)
            p.add_run('\n')
        if data.get('instituicao'):
            r = p.add_run(data['instituicao'])
            set_run_font(r, size=11)
            p.add_run('\n')
        if data.get('ano_conclusao'):
            r = p.add_run(f"Conclusão: {data['ano_conclusao']}")
            set_run_font(r, size=10, color='444444')
        set_paragraph_spacing(p, after=4, line=1.15)

    if data.get('experiencia'):
        add_section_title(doc, 'Experiência profissional')
        if str(data.get('experiencia', '')).strip().lower() in ['sim', '1', 'true', 'yes']:
            add_bullet(doc, 'Possui experiência profissional.')
        else:
            add_bullet(doc, 'Primeiro emprego ou sem experiência formal.')

    if data.get('hard_skills'):
        add_section_title(doc, 'Habilidades técnicas')
        for item in [x.strip() for x in data['hard_skills'].split(',') if x.strip()]:
            add_bullet(doc, item)

    if data.get('soft_skills'):
        add_section_title(doc, 'Habilidades comportamentais')
        for item in [x.strip() for x in data['soft_skills'].split(',') if x.strip()]:
            add_bullet(doc, item)

    if data.get('certificacoes'):
        add_section_title(doc, 'Cursos e certificações')
        p = doc.add_paragraph()
        r = p.add_run(data['certificacoes'])
        set_run_font(r, size=11)
        set_paragraph_spacing(p, after=4, line=1.15)

    if data.get('projetos'):
        add_section_title(doc, 'Projetos')
        p = doc.add_paragraph()
        r = p.add_run(data['projetos'])
        set_run_font(r, size=11)
        set_paragraph_spacing(p, after=4, line=1.15)

    if data.get('observacoes'):
        add_section_title(doc, 'Observações')
        p = doc.add_paragraph()
        r = p.add_run(data['observacoes'])
        set_run_font(r, size=11)
        set_paragraph_spacing(p, after=4, line=1.15)

    doc.save(output_path)
    return output_path


if __name__ == '__main__':
    exemplo = {
        'nome': 'Seu Nome Completo',
        'telefone': '(81) 99999-9999',
        'email': 'seu@email.com',
        'cidade': 'Paulista',
        'estado': 'PE',
        'linkedin': 'linkedin.com/in/seu-perfil',
        'objetivo': 'Atuar na área administrativa, comercial ou tecnologia.',
        'idade': '25',
        'escolaridade': 'Ensino Médio Completo',
        'pretensao': 'A combinar',
        'curso_grad': 'Gestão de Tecnologia da Informação',
        'instituicao': 'Universidade Exemplo',
        'ano_conclusao': '2026',
        'experiencia': 'Não',
        'hard_skills': 'Python, Excel, Word, MQL5, HTML, CSS',
        'soft_skills': 'Comunicação, Organização, Proatividade',
        'certificacoes': 'Curso de Python Básico - 2025',
        'projetos': 'Sistema de cadastro de currículos',
        'observacoes': 'Disponibilidade para início imediato.'
    }
    print(format_docx(exemplo))